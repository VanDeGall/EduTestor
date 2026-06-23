from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import fetch_all

EXPORT_DIR = Path("exports")

PRIORITY_FILLS = {
    "critical": PatternFill("solid", fgColor="FFC7CE"),
    "high": PatternFill("solid", fgColor="F4B183"),
    "medium": PatternFill("solid", fgColor="FFF2CC"),
    "low": PatternFill("solid", fgColor="C6EFCE"),
    "archive": PatternFill("solid", fgColor="D9EAD3"),
}

STATUS_FILLS = {
    "completed": PatternFill("solid", fgColor="C6EFCE"),
    "confirmed": PatternFill("solid", fgColor="C6EFCE"),
    "approved": PatternFill("solid", fgColor="C6EFCE"),
    "waiting": PatternFill("solid", fgColor="D9D2E9"),
    "draft": PatternFill("solid", fgColor="E7E6E6"),
    "archived": PatternFill("solid", fgColor="D9D9D9"),
}


def _write_sheet(ws, headers: list[str], rows: Iterable[dict]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")

    for row in rows:
        values = [row.get(header, "") for header in headers]
        ws.append(values)

    for row in ws.iter_rows(min_row=2):
        header_map = {headers[idx]: cell for idx, cell in enumerate(row)}
        priority = str(header_map.get("priority", "").value or "").lower() if "priority" in header_map else ""
        status = str(header_map.get("status", "").value or "").lower() if "status" in header_map else ""
        fill = PRIORITY_FILLS.get(priority) or STATUS_FILLS.get(status)
        if fill:
            for cell in row:
                cell.fill = fill

    for idx, header in enumerate(headers, start=1):
        max_len = max(len(str(header)), 12)
        for cell in ws[get_column_letter(idx)]:
            if cell.value is not None:
                max_len = max(max_len, min(len(str(cell.value)), 60))
        ws.column_dimensions[get_column_letter(idx)].width = max_len + 2


def _save_workbook(wb: Workbook, filename: str) -> Path:
    EXPORT_DIR.mkdir(exist_ok=True)
    path = EXPORT_DIR / filename
    wb.save(path)
    return path


def export_daily_log(date_value: str) -> Path:
    rows = fetch_all(
        """
        SELECT date, time, clean_activity AS activity, category, sdk_alignment,
               priority, size, output, status, report_include
        FROM activities
        WHERE date = ?
        ORDER BY COALESCE(time, '99:99') ASC, created_at ASC
        """,
        (date_value,),
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily work log"
    _write_sheet(
        ws,
        [
            "date",
            "time",
            "activity",
            "category",
            "sdk_alignment",
            "priority",
            "size",
            "output",
            "status",
            "report_include",
        ],
        rows,
    )
    return _save_workbook(wb, f"{date_value}_sdk_denny-zaznam_v01.xlsx")


def export_monthly_report(month_value: str) -> Path:
    rows = fetch_all(
        """
        SELECT date, activity, category, sdk_alignment, duration_minutes,
               output, evidence_file, status, leadership_relevance
        FROM monthly_sdk_report_view
        WHERE substr(date, 1, 7) = ?
        ORDER BY date ASC
        """,
        (month_value,),
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly SDK report"
    _write_sheet(
        ws,
        [
            "date",
            "activity",
            "category",
            "sdk_alignment",
            "duration_minutes",
            "output",
            "evidence_file",
            "status",
            "leadership_relevance",
        ],
        rows,
    )
    return _save_workbook(wb, f"{month_value}_sdk_mesacny-report_v01.xlsx")


def export_leadership_register(month_value: str | None = None) -> Path:
    if month_value:
        rows = fetch_all(
            """
            SELECT date_prepared, topic, output_type, summary, proposed_decision,
                   status, evidence_file, next_step
            FROM leadership_evidence_view
            WHERE substr(date_prepared, 1, 7) = ?
            ORDER BY date_prepared ASC
            """,
            (month_value,),
        )
        filename = f"{month_value}_vedenie_register-podkladov_v01.xlsx"
    else:
        rows = fetch_all(
            """
            SELECT date_prepared, topic, output_type, summary, proposed_decision,
                   status, evidence_file, next_step
            FROM leadership_evidence_view
            ORDER BY date_prepared ASC
            """
        )
        filename = f"{datetime.now().date()}_vedenie_register-podkladov_v01.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Leadership evidence"
    _write_sheet(
        ws,
        [
            "date_prepared",
            "topic",
            "output_type",
            "summary",
            "proposed_decision",
            "status",
            "evidence_file",
            "next_step",
        ],
        rows,
    )
    return _save_workbook(wb, filename)
